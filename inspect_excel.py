from openpyxl import load_workbook

try:
    wb = load_workbook('BACHIR CODES.xlsx')
    ws = wb.active
    print(f"Sheet Name: {ws.title}")
    
    # Print header
    headers = [cell.value for cell in ws[1]]
    print("Headers:", headers)
    
    # Print first few rows
    for i, row in enumerate(ws.iter_rows(min_row=2, max_row=4, values_only=True), 1):
        print(f"Row {i}: {row}")
        
except Exception as e:
    print(f"Error: {e}")
